import openpyxl

wb = openpyxl.Workbook()                            # Creating an empty workbook

sheets = wb.sheetnames                                # checking what sheets are availables
print('available sheets are: ', str(sheets))

sheet = wb.active
print('checking if the first sheet are accessed: ', sheet.title)


# Changing the title of sheet
sheet.title = 'New Sheet'
print("sheets name after renaming: " , str(wb.sheetnames))


# Creating and Removing sheets
wb.create_sheet()
print("added new sheet: " , str(wb.sheetnames))

wb.create_sheet(index=1, title='First Sheet')
print("added new sheet: " , str(wb.sheetnames))

wb.remove(wb['Sheet'])
print("Workbook after removing sheet: " , str(wb.sheetnames))



# Writing values to cell
sheet['A1'] = 'Hello World!'
print('\nNew added value:', sheet['A1'].value)