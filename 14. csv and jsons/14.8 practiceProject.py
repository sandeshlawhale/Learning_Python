# Project: Excel-to-CSV Converter
# This program finds all Excel files in the current working directory,
# reads every sheet from each workbook,
# and converts each sheet into a separate CSV file.

import openpyxl
import csv
import os

os.chdir('.\\14. csv and jsons')

# Loop through every file in the current working directory
for excelFile in os.listdir('.'):

    # Process only Excel files
    if excelFile.endswith('.xlsx'):

        print('Processing file:', excelFile)

        # Load workbook
        workbook = openpyxl.load_workbook(excelFile)

        # Loop through every sheet in the workbook
        for sheetName in workbook.sheetnames:

            print('Creating CSV file for sheet:', sheetName)

            # Get sheet object
            sheet = workbook[sheetName]

            # Create CSV filename
            csvFileName = excelFile[:-5] + '_' + sheetName + '.csv'

            # Open CSV file in write mode
            csvFile = open(csvFileName, 'w', newline='', encoding='utf-8')

            # Create writer object
            csvWriter = csv.writer(csvFile)

            # Loop through every row in the sheet
            for rowNum in range(1, sheet.max_row + 1):

                # Store row data
                rowData = []

                # Loop through every column in the row
                for colNum in range(1, sheet.max_column + 1):

                    # Get cell value
                    cellValue = sheet.cell(row=rowNum, column=colNum).value

                    # Append cell value to rowData list
                    rowData.append(cellValue)

                # Write row data into CSV file
                csvWriter.writerow(rowData)

            # Close CSV file
            csvFile.close()

print('All Excel files converted to CSV successfully.')