# 12.1 ReadingSpreadsheet.py - read the excel document and process the data

# we will use the openpyxl module for this

import openpyxl, os
from openpyxl.utils import get_column_letter, column_index_from_string
os.chdir('.\\12. Spreadsheets')


# Getting Workbook =========>
wb = openpyxl.load_workbook('example.xlsx')
print('type of the workbook object: ', str(type(wb)))


# Getting sheets from workbook ========>
print('sheets available in the workbook: ') 
sheetNames = wb.sheetnames
print(sheetNames)

print('\nGetting a single Sheet: ')
sheet = wb['Sheet1']
print('type of sheet: ', str(type(sheet)), ', Title of sheet: ', str(sheet.title))

anotherSheet = wb.active
print('title of the active sheet: ', str(anotherSheet.title))


# Getting cells from sheet ========>
value1 = sheet['A1'].value
print('\nvalue of the first cell: ', str(value1))

c = sheet['B1']
print(f'row: {c.row}, column: {c.column}, cords: {c.coordinate}, value: {c.value}')


# Converting between Column letter and Number  ========>
print('\nThe number 1 refers to column: ', get_column_letter(1))
print('The number 2 refers to column: ', get_column_letter(2))

highestCol = sheet.max_column
print('\nThe Last Column of the sheet is: ', get_column_letter(highestCol))

print('\nThe Column Z is at', column_index_from_string('z'), 'th Number')


# get Rows and Column from the sheet ========>
print('\n\nGetting Rows and Column from the sheet')
for Rows in sheet['A1':'C3']:
    for cell in Rows:
        print(cell.coordinate, cell.value)
    print('--------End Of Row--------')


print('\n\nGetting nth column values')
for cell in sheet['B']:
    print(cell.value)

