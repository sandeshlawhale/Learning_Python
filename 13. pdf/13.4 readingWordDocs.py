import docx                                                         # this is useful to interact with word documents, to install this use: pip install python-docx
import os

os.chdir('.\\13. pdf')


doc = docx.Document('demo.docx')
print('length of the docements para: ', len(doc.paragraphs))

print('first para: ', doc.paragraphs[0].text)
print('Second para: ', doc.paragraphs[1].text)

print('\nTotal run in second line: ', doc.paragraphs[1].runs)                               # runs are the continuous texts that has same stylings
print('first run contains: ', doc.paragraphs[1].runs[0].text)
print('second run contains: ', doc.paragraphs[1].runs[1].text)
print('third run contains: ', doc.paragraphs[1].runs[2].text)
print('forth run contains: ', doc.paragraphs[1].runs[3].text)



# getting full docs====>
print('\nGetting Full Docs: ')
def getText(filename):
    doc = docx.Document(filename)
    fullText = []
    for para in doc.paragraphs:
        fullText.append(para.text)
    return '\n'.join(fullText)

print(getText('demo.docx'))



# stylings =>>>>>
doc.paragraphs[0].style = 'Normal'
doc.paragraphs[1].runs[0].style = 'QuoteChar'


# writing word docx
newDoc = docx.Document()
newDoc.add_paragraph('hello World!', 'Title')                           # we passed the (text, style), we can only pass the text also
paraObj1 = newDoc.add_paragraph('new para for new start.')
paraObj1.add_run('adding run to test.')
newDoc.add_page_break()
newDoc.add_paragraph('this should be on second page.')
newDoc.save('helloWorld.docx')